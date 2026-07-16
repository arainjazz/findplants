#!/usr/bin/env python3
# Generate 6 provincial protected-list seed migrations from the user's workbook.
# Reads 总表 (all taxa) + 来源 (per-province source metadata).
# 条目类型: 种级条目 -> species ; 属级条目 -> genus. status='省级' (provincial lists have no 一级/二级).
import openpyxl, glob, re, unicodedata

XLSX = glob.glob("/sessions/affectionate-gracious-hopper/mnt/uploads/*.xlsx")[0]
OUTDIR = "/sessions/affectionate-gracious-hopper/mnt/plantspedia/supabase/migrations/"

def fold(s): return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
def norm(s):
    # Mirror of JS normalizeSciName (src/lib/catalogs.ts). The hybrid marker (× or a
    # standalone "x") is dropped: "Populus × irtyschensis" -> "populus irtyschensis".
    s = fold(s or ""); s = re.sub(r"[*_]", "", s); s = re.sub(r"\([^)]*\)", " ", s)
    s = re.sub(r"[×✕⨯]", " ", s)
    s = re.sub(r"\s+", " ", s).strip().lower(); t = [x for x in s.split(" ") if x and x != "x"]
    return f"{t[0]} {t[1]}" if len(t) >= 2 else (t[0] if t else "")

# province -> (filename, listName, version, dateISO)
META = {
    "云南":   ("20260703150000_seed_yunnan_2023.sql",    "云南（2023）",  "2023", "2023-12-15"),
    "内蒙古": ("20260703160000_seed_neimenggu_2009.sql", "内蒙古（2009）", "2009", "2009-07-30"),
    "四川":   ("20260703170000_seed_sichuan_2024.sql",   "四川（2024）",  "2024", "2024-08-05"),
    "广东":   ("20260703180000_seed_guangdong_2023.sql", "广东（2023）",  "2023", "2023-03-17"),
    "贵州":   ("20260703190000_seed_guizhou_2023.sql",   "贵州（2023）",  "2023", "2023-11-28"),
    "福建":   ("20260703200000_seed_fujian_2024.sql",    "福建（2024）",  "2024", "2024-01-29"),
}

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

# source notes from 来源
src = {}
ws = wb['来源']; rows = list(ws.iter_rows(values_only=True)); hdr = rows[3]
for r in rows[4:]:
    if not r[0]: continue
    d = dict(zip(hdr, r))
    prov = d['省区']
    note = f"{d['名录名称']}｜{d['文件文号']}｜{str(d['发布日期'])[:10]} 发布｜{d['范围与说明']}"
    src[prov] = (note, d['来源网址'])

# taxa from 总表
ws = wb['总表']; rows = list(ws.iter_rows(values_only=True)); hdr = rows[2]
data = [dict(zip(hdr, r)) for r in rows[3:] if r[0]]

def sq(s): return "'" + str(s).replace("'", "''") + "'"

for prov, (fname, lname, ver, date) in META.items():
    note, url = src[prov]
    prov_rows = [d for d in data if d['省区'] == prov]
    seen = set(); taxa = []
    for d in prov_rows:
        sci = (d['学名'] or "").strip()
        zh = (d['中文名'] or "").strip().rstrip("*").strip()
        if not sci: continue
        if d['条目类型'] and '属级' in str(d['条目类型']):
            rank = 'genus'; g = fold(sci.split()[0]).lower(); n = g
        else:
            rank = 'species'; n = norm(sci)
        key = (rank, n)
        if key in seen: continue
        seen.add(key)
        taxa.append((sci, n, zh, rank))
    o = []
    o.append(f"-- Seed: {lname} 省级重点保护野生植物名录. province={prov}. {len(taxa)} taxa.")
    o.append(f"-- Source: {note}")
    o.append("DELETE FROM public.conservation_taxa t USING public.conservation_lists l")
    o.append(f"  WHERE t.list_id = l.id AND l.kind = 'protected' AND l.province = {sq(prov)};")
    o.append(f"DELETE FROM public.conservation_lists WHERE kind = 'protected' AND province = {sq(prov)};")
    o.append("WITH new_list AS (")
    o.append("  INSERT INTO public.conservation_lists (kind, name, province, version, effective_date, source_url, source_note)")
    o.append(f"  VALUES ('protected', {sq(lname)}, {sq(prov)}, {sq(ver)}, DATE {sq(date)}, {sq(url)}, {sq(note)})")
    o.append("  RETURNING id")
    o.append(")")
    o.append("INSERT INTO public.conservation_taxa (list_id, scientific_name, normalized_name, chinese_name, status, rank, excluded_names)")
    o.append("SELECT nl.id, v.sci, v.norm, v.zh, '省级', v.rank, NULL::text[] FROM new_list nl, (VALUES")
    vals = [f"  ({sq(sci)}, {sq(n)}, {sq(zh)}, {sq(rank)})" for sci, n, zh, rank in taxa]
    o.append(",\n".join(vals))
    o.append(") AS v(sci, norm, zh, rank);")
    open(OUTDIR + fname, "w", encoding="utf-8").write("\n".join(o) + "\n")
    print(f"{prov}: {len(taxa)} taxa -> {fname}")
PYEOF_MARKER = None
